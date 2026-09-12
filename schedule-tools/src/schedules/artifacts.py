from __future__ import annotations

import hashlib
import json
import re
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from .paths import DATA_DIR, all_review_dirs, artifact_path, parse_review_dir_name, relative_to_repo


class PrefixCollisionError(RuntimeError):
    """Two review dirs share a sha12 prefix but differ in full hash."""


def _sidecar_sha256(review_dir: Path) -> str | None:
    sidecar = review_dir / "source.sha256"
    if not sidecar.is_file():
        return None
    try:
        text = sidecar.read_text().strip()
    except OSError:
        return None
    if len(text) == 64 and all(char in "0123456789abcdef" for char in text):
        return text
    return None


def _hash_and_backfill_source_pdf(review_dir: Path) -> str | None:
    pdf = review_dir / "source.pdf"
    if not pdf.is_file():
        return None
    digest = hashlib.sha256(pdf.read_bytes()).hexdigest()
    (review_dir / "source.sha256").write_text(f"{digest}\n")
    return digest


def find_review_dir_for_sha(slug: str, sha256: str, *, root: Path = DATA_DIR) -> Path | None:
    """Return the review dir whose full source hash is ``sha256``.

    Only parsed ``<date>-<sha12>`` dirs whose sha12 matches ``sha256[:12]``
    are considered. A present ``source.sha256`` sidecar wins; otherwise the
    helper hashes ``source.pdf`` and backfills the sidecar. Any prefix match
    with a different full hash is a collision.
    """
    prefix = sha256[:12]
    hits: list[Path] = []
    collision: PrefixCollisionError | None = None
    for review_dir in all_review_dirs(slug, root=root):
        parsed = parse_review_dir_name(review_dir.name)
        if parsed is None or parsed[1] != prefix:
            continue
        full = _sidecar_sha256(review_dir)
        if full is None:
            full = _hash_and_backfill_source_pdf(review_dir)
        if full is None:
            continue
        if full != sha256:
            collision = PrefixCollisionError(
                f"prefix collision in {slug}: existing={full} new={sha256}"
            )
            continue
        hits.append(review_dir)
    if collision is not None:
        raise collision
    return hits[0] if hits else None


def _sha256_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _sha256_json(value: dict) -> str:
    return hashlib.sha256(json.dumps(value, sort_keys=True).encode("utf-8")).hexdigest()


# Cloudflare rewrites its email obfuscation on every response, so these two
# payloads carry no schedule and must not change a capture's identity.
_CFEMAIL_ATTRIBUTE = re.compile(rb"""\s*data-cfemail=["'][0-9a-fA-F]*["']""")
_CFEMAIL_HREF = re.compile(rb"/cdn-cgi/l/email-protection#[0-9a-fA-F]*")


def workbook_facts(content: bytes) -> dict[str, dict]:
    """Everything the workbook parser reads: cell values, visibility, merges.

    The Google Sheets export reorders ``sharedStrings.xml`` and ``styles.xml``
    between downloads of an unchanged sheet, so the archive bytes are not an
    identity. What the parser reads is: values, which sheets are visible
    (``providers/koret.py`` skips the hidden ones), and the merged ranges it
    reads closures and banners from.
    """
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(content), data_only=True)
    return {
        sheet.title: {
            "state": sheet.sheet_state,
            "merges": sorted(str(area) for area in sheet.merged_cells.ranges),
            "cells": [[None if cell.value is None else str(cell.value) for cell in row]
                      for row in sheet.iter_rows()],
        }
        for sheet in workbook.worksheets
    }


def canonical_source_sha256(kind: str, content: bytes) -> str:
    """The identity of a captured document, ignoring per-response noise.

    Raw bytes stay the published identity: ``source.sha256`` and every
    artifact's ``pdf_sha256`` keep hashing the file as it was stored. This hash
    answers a narrower question — is this fetch the document we already hold? —
    so a rotating Cloudflare token or a reshuffled spreadsheet archive cannot
    mint a new snapshot for an unchanged schedule.
    """
    if kind == "html":
        content = _CFEMAIL_HREF.sub(b"/cdn-cgi/l/email-protection",
                                    _CFEMAIL_ATTRIBUTE.sub(b"", content))
    elif kind == "xlsx":
        content = json.dumps(workbook_facts(content), sort_keys=True).encode("utf-8")
    return hashlib.sha256(content).hexdigest()


def save_artifact_bundle(
    *,
    slug: str,
    date: str,
    provider: str,
    model: str,
    source_pdf_url: str,
    pdf_sha256: str,
    prompt: str,
    schema: dict,
    payload: dict,
    usage: dict,
    cost_estimate: str,
    details: dict | None = None,
    root: Path = DATA_DIR,
) -> dict[str, str]:
    target = artifact_path(slug, date, pdf_sha256, provider, model, root=root)
    target.parent.mkdir(parents=True, exist_ok=True)

    provider_payload: dict = {
        "provider": provider,
        "model": model,
        "extracted_at": datetime.now(timezone.utc).isoformat(),
        "prompt_sha256": _sha256_text(prompt),
        "schema_sha256": _sha256_json(schema),
        "source_pdf_url": source_pdf_url,
        "pdf_sha256": pdf_sha256,
        "usage": usage,
        "cost_estimate": cost_estimate,
        "payload": payload,
    }
    if details is not None:
        provider_payload["details"] = details
    target.write_text(json.dumps(provider_payload, indent=2, sort_keys=True) + "\n")

    return {provider: relative_to_repo(target)}


def skip_if_fresh(
    *,
    slug: str,
    date: str,
    pdf_sha256: str,
    provider: str,
    model: str,
    prompt: str,
    schema: dict,
    configuration: dict | None = None,
    root: Path = DATA_DIR,
) -> bool:
    """Return True iff a cached provider JSON exists and its hashes match."""
    provider_file = artifact_path(slug, date, pdf_sha256, provider, model, root=root)
    if not provider_file.exists():
        return False
    try:
        data = json.loads(provider_file.read_text())
    except (OSError, json.JSONDecodeError):
        return False
    return (
        data.get("prompt_sha256") == _sha256_text(prompt)
        and data.get("schema_sha256") == _sha256_json(schema)
        and (configuration is None or data.get("details", {}).get("configuration") == configuration)
    )


def pool_bundle_identity(sources: list[dict]) -> str:
    return _sha256_json({"kind": "north_beach_cool_warm", "sources": [{key: value for key, value in source.items() if key != "capture"} for source in sources]})


def combine_pool_artifacts(sources: list[dict], artifacts: list[dict], documents: list[bytes], prompt: str) -> dict:
    from .providers.openai_provider import verify_artifact, ClosureReviewRequired
    from .signals import inspect_pdf_source, north_beach_pool_identity
    from .grounding import source_slots
    from .validate import validate

    if len(sources) != 2 or [item["pool"] for item in sources] != ["cool", "warm"]:
        raise ValueError("A bundle requires exactly one Cool and one Warm source")
    if len({item["url"] for item in sources}) != 2 or len({item["sha256"] for item in sources}) != 2:
        raise ValueError("Duplicate pool source")
    windows = []
    sessions = []
    facility_closures = []
    scoped_closures = []
    for member, artifact, document in zip(sources, artifacts, documents, strict=True):
        source = inspect_pdf_source(document)
        if north_beach_pool_identity(source.text) != member["pool"]:
            raise ValueError("Printed physical pool differs from the bundle")
        if artifact.get("source_pdf_url") != member["url"] or artifact.get("pdf_sha256") != member["sha256"]:
            raise ValueError("Component source identity differs from the bundle")
        if artifact.get("details", {}).get("configuration") != member["configuration"]:
            raise ValueError("Component configuration differs from the bundle")
        coverage = verify_artifact(artifact, document, prompt)
        if not coverage["ok"] or not validate(artifact["payload"]).ok:
            raise ValueError("Component failed independent source verification")
        payload = artifact["payload"]
        if payload.get("schedule_basis") != "swim_schedule" or payload.get("access_hours") or payload.get("access_exceptions"):
            raise ValueError("Unsupported pool component payload")
        windows.append((payload["effective_start"], payload.get("effective_end")))
        slots = {slot.key: slot.cell.id for slot in source_slots(source)}
        facts = artifact["details"]["source_facts"]["sessions"]
        for session, fact in zip(payload["sessions"], facts, strict=True):
            key = tuple(session.get(field) for field in ("day", "type", "start", "end", "pool"))
            sessions.append(session | {"physical_pool": member["pool"], "pool_label_raw": fact["pool_label_raw"],
                                       "source_sha256": member["sha256"], "source_cell": slots[key]})
        facility_closures.append([closure for closure in payload["closures"] if not closure.get("physical_pool")])
        scoped_closures.extend(closure | {"source_notices": [notice | {"source_sha256": member["sha256"]}
                                                           for notice in closure["source_notices"]]}
                               for closure in payload["closures"] if closure.get("physical_pool"))
    if windows[0] != windows[1] or not windows[0][1]:
        raise ValueError("Pool effective windows conflict")
    closure_key = lambda closure: tuple(closure.get(field, "") for field in ("start", "end", "start_time", "end_time", "reason_code"))
    if sorted(map(closure_key, facility_closures[0])) != sorted(map(closure_key, facility_closures[1])):
        raise ClosureReviewRequired(source, ["pair:facility_closure_conflict"])
    combined_closures = []
    for closure in facility_closures[0]:
        other = next(item for item in facility_closures[1] if closure_key(item) == closure_key(closure))
        notices = [notice | {"source_sha256": sources[index]["sha256"]}
                   for index, item in enumerate((closure, other)) for notice in item["source_notices"]]
        combined_closures.append(closure | {"source_notices": notices})
    payload = {"effective_start": windows[0][0], "effective_end": windows[0][1], "schedule_basis": "swim_schedule",
               "sessions": sessions, "closures": combined_closures + scoped_closures}
    if not validate(payload).ok:
        raise ValueError("Invalid combined pool schedule")
    return payload


def save_pool_bundle(slug: str, component_paths: list[Path], prompt: str) -> Path:
    from .providers.openai_provider import API_MODEL
    from .signals import inspect_pdf_source, north_beach_pool_identity
    from .paths import review_dir
    if slug != "north-beach-pool" or len(component_paths) != 2:
        raise ValueError("Unsupported source bundle")
    components = []
    for path in component_paths:
        artifact = json.loads(path.read_text())
        document = (path.parent / "source.pdf").read_bytes()
        pool = north_beach_pool_identity(inspect_pdf_source(document).text)
        member = {"pool": pool, "url": artifact["source_pdf_url"], "sha256": artifact["pdf_sha256"],
                  "capture": path.parent.name, "configuration": artifact["details"]["configuration"]}
        components.append((member, artifact, document))
    components.sort(key=lambda item: item[0]["pool"] or "")
    sources, artifacts, documents = map(list, zip(*components, strict=True))
    payload = combine_pool_artifacts(sources, artifacts, documents, prompt)
    identity = pool_bundle_identity(sources)
    root = component_paths[0].parent.parent.parent
    fetch_date = max(item["capture"][:10] for item in sources)
    directory = review_dir(slug, fetch_date, identity, root=root)
    directory.mkdir(parents=True, exist_ok=True)
    manifest = directory / "source-bundle.json"
    text = json.dumps(sources, indent=2) + "\n"
    if manifest.exists() and manifest.read_text() != text:
        raise ValueError("Bundle identity collision")
    manifest.write_text(text)
    target = directory / "openai-pool-bundle.json"
    bundle = {"provider": "openai", "model": API_MODEL, "bundle_sha256": identity, "source_bundle": sources,
              "payload": payload, "extracted_at": max(item["extracted_at"] for item in artifacts)}
    target.write_text(json.dumps(bundle, indent=2) + "\n")
    return target


def verify_pool_bundle(artifact: dict, directory: Path, prompt: str) -> dict:
    from .paths import slugify
    from .providers.openai_provider import API_MODEL
    sources = artifact["source_bundle"]
    if artifact.get("provider") != "openai" or artifact.get("model") != API_MODEL or directory.parent.name != "north-beach-pool":
        raise ValueError("Unsupported bundle provider or facility")
    identity = pool_bundle_identity(sources)
    if artifact.get("bundle_sha256") != identity or directory.name[11:] != identity[:12]:
        raise ValueError("Bundle identity mismatch")
    if json.loads((directory / "source-bundle.json").read_text()) != sources:
        raise ValueError("Bundle manifest mismatch")
    artifacts, documents = [], []
    for member in sources:
        capture = member["capture"]
        if not parse_review_dir_name(capture) or capture[11:] != member["sha256"][:12]:
            raise ValueError("Invalid component capture")
        parent = directory.parent / capture
        artifacts.append(json.loads((parent / f"openai-{slugify(API_MODEL)}.json").read_text()))
        documents.append((parent / "source.pdf").read_bytes())
    expected = combine_pool_artifacts(sources, artifacts, documents, prompt)
    if artifact.get("payload") != expected:
        raise ValueError("Bundle payload differs from its verified components")
    return {"ok": True}
